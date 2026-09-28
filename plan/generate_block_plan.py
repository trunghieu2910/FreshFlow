import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

def build_block_plan():
    wb_path = 'd:/FreshFlow/plan/backlog-freshflow-mvp-12-tuan-updated.xlsx'
    wb = openpyxl.load_workbook(wb_path)
    source_sheet = wb['Backlog 12 Tuan']

    # 1. Read all tasks from Backlog 12 Tuan
    tasks_dict = {}
    for r in range(3, source_sheet.max_row + 1):
        tid = source_sheet.cell(row=r, column=1).value
        if not tid:
            continue
        row_vals = [source_sheet.cell(row=r, column=c).value for c in range(1, 22)]
        tasks_dict[tid] = {
            'task_id': row_vals[0],
            'orig_week': row_vals[1],
            'orig_theme': row_vals[2],
            'orig_day': row_vals[3],
            'orig_task_idx': row_vals[4],
            'track': row_vals[5],
            'name': row_vals[6],
            'goal': row_vals[7],
            'how': row_vals[8],
            'output': row_vals[9],
            'criteria': row_vals[10],
            'test': row_vals[11],
            'est': row_vals[12],
            'priority': row_vals[13],
            'dependencies': row_vals[14],
            'status': row_vals[15],
            'evidence': row_vals[16],
            'note': row_vals[17],
            'db_task_id': row_vals[18],
            'db_focus': row_vals[19],
            'db_output': row_vals[20],
        }

    print(f"Loaded {len(tasks_dict)} tasks from source sheet.")

    # 2. Define Blocks (Cục) and their Sub-blocks with exact task IDs
    blocks = [
        {
            'block_id': 'CỤC 1',
            'block_name': 'CỤC 1: NỀN TẢNG HỆ THỐNG & MÔ HÌNH HÓA DOMAIN (Foundation & Architecture)',
            'block_desc': 'Thiết lập monorepo, quy ước kiến trúc, thiết kế domain model, state machine, docker database và ERD nền móng',
            'subblocks': [
                {
                    'sub_name': '1.1. Phạm vi MVP & Thiết kế Domain State Machine',
                    'task_ids': ['FF-01-01-1', 'FF-01-01-2', 'FF-01-02-1', 'FF-01-02-2']
                },
                {
                    'sub_name': '1.2. Khởi tạo Backend Skeleton & Hạ tầng PostgreSQL Docker',
                    'task_ids': ['FF-01-03-1', 'FF-01-03-2', 'FF-01-04-1', 'FF-01-04-2']
                },
                {
                    'sub_name': '1.3. Mô hình dữ liệu ERD, Quy trình Backlog & Quản trị rủi ro',
                    'task_ids': ['FF-01-05-1', 'FF-01-05-2', 'FF-01-06-1', 'FF-01-06-2', 'FF-01-07-1', 'FF-01-07-2']
                }
            ]
        },
        {
            'block_id': 'CỤC 2',
            'block_name': 'CỤC 2: TOÀN BỘ BACKEND & HỆ THỐNG APIS (Complete Backend Services, DB & REST)',
            'block_desc': 'Xây dựng trọn vẹn toàn bộ Backend từ Catalog, Auth/RBAC, Order State Machine, Inventory Concurrency, Payment Mock đến Delivery Driver API trước khi chuyển sang Client',
            'subblocks': [
                {
                    'sub_name': '2.1. Catalog & Store Module (JPA Entities, Flyway, CRUD, Search & Pagination)',
                    'task_ids': [
                        'FF-02-01-1', 'FF-02-01-2', 'FF-02-02-1', 'FF-02-02-2',
                        'FF-02-03-1', 'FF-02-03-2', 'FF-02-04-1', 'FF-02-04-2',
                        'FF-02-05-1', 'FF-02-05-2', 'FF-02-06-1', 'FF-02-06-2',
                        'FF-02-07-1', 'FF-02-07-2'
                    ]
                },
                {
                    'sub_name': '2.2. User, Store Ownership & Authentication/RBAC Core (Spring Security & JWT)',
                    'task_ids': [
                        'FF-04-01-2', 'FF-07-01-1', 'FF-07-01-2', 'FF-07-02-1',
                        'FF-07-02-2', 'FF-07-05-1', 'FF-07-05-2', 'FF-07-06-2',
                        'FF-07-07-1', 'FF-07-07-2'
                    ]
                },
                {
                    'sub_name': '2.3. Customer Public Catalog Endpoint (Public API for Clients)',
                    'task_ids': ['FF-05-05-2']
                },
                {
                    'sub_name': '2.4. Order Management & Order State Machine Core (Create, Transitions & History)',
                    'task_ids': [
                        'FF-06-03-1', 'FF-06-03-2', 'FF-06-05-1',
                        'FF-04-02-2', 'FF-04-03-2', 'FF-04-04-2', 'FF-04-05-2'
                    ]
                },
                {
                    'sub_name': '2.5. Inventory, Daily Capacity & Idempotent Checkout (Concurrency & Audit)',
                    'task_ids': [
                        'FF-06-07-1', 'FF-08-01-1', 'FF-08-01-2', 'FF-08-02-1',
                        'FF-08-02-2', 'FF-08-03-1', 'FF-08-04-2', 'FF-08-05-1',
                        'FF-08-05-2', 'FF-08-06-1', 'FF-08-06-2', 'FF-08-07-1', 'FF-08-07-2'
                    ]
                },
                {
                    'sub_name': '2.6. Payment Mock, Acceptance Compensation & Delivery/Driver APIs',
                    'task_ids': [
                        'FF-09-01-1', 'FF-09-02-1', 'FF-09-03-1', 'FF-09-03-2',
                        'FF-09-04-1', 'FF-09-05-1', 'FF-09-07-1'
                    ]
                },
                {
                    'sub_name': '2.7. Backend Optimization, Integration Tests & Technical Hardening',
                    'task_ids': [
                        'FF-10-01-1', 'FF-10-01-2', 'FF-10-04-1',
                        'FF-11-02-1', 'FF-11-02-2', 'FF-11-03-1', 'FF-11-03-2'
                    ]
                }
            ]
        },
        {
            'block_id': 'CỤC 3',
            'block_name': 'CỤC 3: FRONTEND - MOBILE CLIENT (Android Kotlin: Customer & Driver App)',
            'block_desc': 'Phát triển ứng dụng Android native bằng Jetpack Compose: luồng khách hàng (Catalog, Giỏ hàng offline, Checkout, Theo dõi đơn) và luồng tài xế (Bật nhận đơn, Nhận đơn, OTP/COD, Báo lỗi)',
            'subblocks': [
                {
                    'sub_name': '3.1. Nền tảng Android App, Architecture MVVM & Design System',
                    'task_ids': ['FF-04-07-2', 'FF-05-01-1', 'FF-05-01-2', 'FF-05-02-1']
                },
                {
                    'sub_name': '3.2. Customer Catalog Browsing & Detail Screens',
                    'task_ids': [
                        'FF-05-02-2', 'FF-05-03-1', 'FF-05-03-2', 'FF-05-04-1',
                        'FF-05-04-2', 'FF-05-05-1', 'FF-05-06-1', 'FF-05-06-2',
                        'FF-05-07-1', 'FF-05-07-2'
                    ]
                },
                {
                    'sub_name': '3.3. Offline Cart Persistence với Room Database & CartRepository',
                    'task_ids': ['FF-06-01-1', 'FF-06-01-2', 'FF-06-02-1']
                },
                {
                    'sub_name': '3.4. Mobile Authentication & DataStore Token Management',
                    'task_ids': ['FF-06-07-2', 'FF-07-04-1', 'FF-07-04-2']
                },
                {
                    'sub_name': '3.5. Customer Checkout, Stock Invalidation & Order Tracking UI',
                    'task_ids': [
                        'FF-06-04-1', 'FF-06-04-2', 'FF-06-05-2',
                        'FF-06-06-1', 'FF-06-06-2', 'FF-08-04-1', 'FF-09-02-2'
                    ]
                },
                {
                    'sub_name': '3.6. Mobile Driver Client Workflow (Nhận đơn, OTP/COD & Delivery Report)',
                    'task_ids': ['FF-09-04-2', 'FF-09-05-2', 'FF-09-06-2']
                },
                {
                    'sub_name': '3.7. Mobile Quality, Unit Tests & Polish Lifecycle/Rotation',
                    'task_ids': ['FF-10-02-2', 'FF-11-01-2', 'FF-11-04-2']
                }
            ]
        },
        {
            'block_id': 'CỤC 4',
            'block_name': 'CỤC 4: FRONTEND - WEB MERCHANT (React TypeScript: Catalog, Order Dashboard & Operations)',
            'block_desc': 'Phát triển cổng thông tin quản lý cho chủ cửa hàng (Merchant Portal): Quản lý danh mục, Biến thể món, Sức chứa hàng ngày, Tiếp nhận & Xử lý đơn hàng, Gán tài xế và Khiếu nại',
            'subblocks': [
                {
                    'sub_name': '4.1. Web Setup, AppShell Layout, Router & API Client',
                    'task_ids': ['FF-03-01-1', 'FF-03-01-2', 'FF-03-02-1', 'FF-03-02-2']
                },
                {
                    'sub_name': '4.2. Merchant Catalog Management (CRUD Product, Variant & Capacity)',
                    'task_ids': [
                        'FF-03-03-1', 'FF-03-03-2', 'FF-03-04-1', 'FF-03-04-2',
                        'FF-03-05-1', 'FF-03-05-2', 'FF-03-06-1', 'FF-03-06-2',
                        'FF-03-07-1', 'FF-03-07-2'
                    ]
                },
                {
                    'sub_name': '4.3. Merchant Order Dashboard & State Transitions (Accept/Reject/Dispatch)',
                    'task_ids': [
                        'FF-04-01-1', 'FF-04-02-1', 'FF-04-03-1', 'FF-04-04-1',
                        'FF-04-05-1', 'FF-04-06-1', 'FF-04-06-2', 'FF-04-07-1'
                    ]
                },
                {
                    'sub_name': '4.4. Web Customer Preview & Auth Integration (AuthContext, Interceptor)',
                    'task_ids': ['FF-06-02-2', 'FF-07-03-1', 'FF-07-03-2']
                },
                {
                    'sub_name': '4.5. Merchant Inventory Adjustment & Order Workflow Integration UI',
                    'task_ids': ['FF-08-03-2', 'FF-09-01-2']
                },
                {
                    'sub_name': '4.6. Web Component Tests, Accessibility (a11y) & UX Polish',
                    'task_ids': ['FF-10-02-1', 'FF-11-01-1', 'FF-11-04-1']
                }
            ]
        },
        {
            'block_id': 'CỤC 5',
            'block_name': 'CỤC 5: TÍCH HỢP HỆ THỐNG ĐA NỀN TẢNG, DEVOPS & E2E TESTING (Full Integration & QA)',
            'block_desc': 'Hợp nhất 3 client (Backend + Web + Mobile), triển khai Docker Compose toàn diện, kiểm thử E2E kịch bản giao dịch thực tế và diễn tập triển khai',
            'subblocks': [
                {
                    'sub_name': '5.1. Tích hợp xác thực & Demo luồng xuyên suốt Customer-Merchant-Driver',
                    'task_ids': ['FF-07-06-1', 'FF-09-06-1', 'FF-09-07-2']
                },
                {
                    'sub_name': '5.2. DevOps, Docker Compose Full Stack & Architecture Diagrams',
                    'task_ids': ['FF-10-03-1', 'FF-10-03-2', 'FF-10-04-2']
                },
                {
                    'sub_name': '5.3. Bảo mật, Hardening Config, Clean Deployment Rehearsal & Scope Freeze',
                    'task_ids': ['FF-10-05-1', 'FF-10-05-2', 'FF-10-06-1', 'FF-10-06-2', 'FF-10-07-1', 'FF-10-07-2']
                },
                {
                    'sub_name': '5.4. Full Regression Demo & Dọn dẹp mã nguồn cuối kỳ',
                    'task_ids': ['FF-11-06-1', 'FF-11-06-2']
                }
            ]
        },
        {
            'block_id': 'CỤC 6',
            'block_name': 'CỤC 6: ĐÓNG GÓI SẢN PHẨM, DEMO & HỒ SƠ NGHỀ NGHIỆP (Release v1.0.0, Portfolio & Career)',
            'block_desc': 'Đóng gói bản phát hành chính thức v1.0.0, quay video demo hoàn chỉnh, xây dựng portfolio dự án, case study kỹ thuật và chuẩn bị hồ sơ ứng tuyển',
            'subblocks': [
                {
                    'sub_name': '6.1. Thu thập Media, Video Demo 3 Client & Hồ sơ Portfolio v1',
                    'task_ids': ['FF-11-05-1', 'FF-11-05-2', 'FF-11-07-1', 'FF-11-07-2']
                },
                {
                    'sub_name': '6.2. Sửa lỗi phát hành cuối cùng, Tài liệu hoá & Gắn tag Release v1.0.0',
                    'task_ids': ['FF-12-01-1', 'FF-12-01-2', 'FF-12-02-1', 'FF-12-02-2']
                },
                {
                    'sub_name': '6.3. Kịch bản Demo 10-15 phút & Dry Run hoàn hảo',
                    'task_ids': ['FF-12-03-1', 'FF-12-03-2']
                },
                {
                    'sub_name': '6.4. GitHub Presentation, Project Case Study & Nghề nghiệp / CV',
                    'task_ids': ['FF-12-04-1', 'FF-12-04-2', 'FF-12-05-1', 'FF-12-05-2']
                },
                {
                    'sub_name': '6.5. Final Retrospective 3 Client, Đóng gói Archive & Bắt đầu Ứng tuyển',
                    'task_ids': ['FF-12-06-1', 'FF-12-06-2', 'FF-12-07-1', 'FF-12-07-2']
                }
            ]
        }
    ]

    # Verify counts
    total_tasks_in_blocks = 0
    assigned_tids = set()
    for b in blocks:
        b_count = 0
        for sb in b['subblocks']:
            b_count += len(sb['task_ids'])
            for tid in sb['task_ids']:
                assert tid in tasks_dict, f"Unknown task {tid}"
                assert tid not in assigned_tids, f"Duplicate task {tid}"
                assigned_tids.add(tid)
        print(f"{b['block_id']}: {b_count} tasks")
        total_tasks_in_blocks += b_count

    assert total_tasks_in_blocks == 168, f"Expected 168 tasks, got {total_tasks_in_blocks}"
    assert len(assigned_tids) == 168
    print("ALL 168 tasks perfectly verified across 6 blocks!")

    # 3. Create or replace new worksheet
    new_sheet_name = 'Kế hoạch theo cục'
    if new_sheet_name in wb.sheetnames:
        del wb[new_sheet_name]

    ws = wb.create_sheet(title=new_sheet_name)

    # Styles
    font_family = 'Calibri'
    title_font = Font(name=font_family, size=16, bold=True, color='FFFFFF')
    title_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')

    header_font = Font(name=font_family, size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='17365D', end_color='17365D', fill_type='solid')

    block_banner_font = Font(name=font_family, size=12, bold=True, color='FFFFFF')
    block_banner_fill = PatternFill(start_color='203764', end_color='203764', fill_type='solid')

    subblock_banner_font = Font(name=font_family, size=11, bold=True, color='1F4E79')
    subblock_banner_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

    regular_font = Font(name=font_family, size=10)
    bold_font = Font(name=font_family, size=10, bold=True)

    # Status styles
    status_styles = {
        'Done': {
            'font': Font(name=font_family, size=10, bold=True, color='276A3C'),
            'fill': PatternFill(start_color='D9EAD3', end_color='D9EAD3', fill_type='solid')
        },
        'Review': {
            'font': Font(name=font_family, size=10, bold=True, color='843B62'),
            'fill': PatternFill(start_color='FCE5CD', end_color='FCE5CD', fill_type='solid')
        },
        'In Progress': {
            'font': Font(name=font_family, size=10, bold=True, color='8F5B00'),
            'fill': PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
        },
        'Todo': {
            'font': Font(name=font_family, size=10, color='595959'),
            'fill': PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
        }
    }

    thin_border_side = Side(border_style='thin', color='D9D9D9')
    table_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Title row
    ws.merge_cells('A1:U1')
    title_cell = ws['A1']
    title_cell.value = "FRESHFLOW MVP — KẾ HOẠCH PHÁT TRIỂN THEO CỤC (BACKEND FIRST → FRONTEND: MOBILE & REACT → INTEGRATION & RELEASE)"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 40

    # Header Row
    headers = [
        'STT Mới',              # Col 1
        'Task ID',              # Col 2
        'Cục thực hiện',        # Col 3
        'Phân nhóm chức năng',  # Col 4
        'Tuần gốc',             # Col 5
        'Track',                # Col 6
        'Task cần làm',         # Col 7
        'Mục tiêu học',         # Col 8
        'Cách thực hiện',       # Col 9
        'Đầu ra bắt buộc',      # Col 10
        'Acceptance criteria',  # Col 11
        'Test/kiểm chứng',      # Col 12
        'Ước tính (giờ)',       # Col 13
        'Ưu tiên',              # Col 14
        'Phụ thuộc',            # Col 15
        'Trạng thái (Giữ nguyên)', # Col 16
        'Link bằng chứng',      # Col 17
        'Ghi chú sắp xếp',      # Col 18
        'Database Task ID',     # Col 19
        'Database focus',       # Col 20
        'Database output'       # Col 21
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = table_border
    ws.row_dimensions[2].height = 32

    current_row = 3
    global_seq = 1

    for block in blocks:
        # Block separator row
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
        b_cell = ws.cell(row=current_row, column=1)
        b_cell.value = f"❖ {block['block_name']} — {block['block_desc']}"
        b_cell.font = block_banner_font
        b_cell.fill = block_banner_fill
        b_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.row_dimensions[current_row].height = 28
        current_row += 1

        for subblock in block['subblocks']:
            # Sub-block separator row
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
            sb_cell = ws.cell(row=current_row, column=1)
            sb_cell.value = f"  ▶ {subblock['sub_name']} ({len(subblock['task_ids'])} tasks)"
            sb_cell.font = subblock_banner_font
            sb_cell.fill = subblock_banner_fill
            sb_cell.alignment = Alignment(horizontal='left', vertical='center', indent=2)
            ws.row_dimensions[current_row].height = 24
            current_row += 1

            for tid in subblock['task_ids']:
                t = tasks_dict[tid]

                # Rationale/Note enhancement
                orig_note = str(t['note']) if t['note'] else ''
                rationale_note = orig_note
                if block['block_id'] == 'CỤC 2':
                    prefix = "[Backend-First] "
                    if not rationale_note.startswith(prefix):
                        rationale_note = prefix + (orig_note if orig_note else "Hoàn thành toàn bộ API & DB trước khi chuyển sang client.")
                elif block['block_id'] == 'CỤC 3':
                    prefix = "[Mobile Frontend] "
                    if not rationale_note.startswith(prefix):
                        rationale_note = prefix + (orig_note if orig_note else "Xây dựng native Android app (Customer & Driver) trên nền Backend hoàn chỉnh.")
                elif block['block_id'] == 'CỤC 4':
                    prefix = "[React Web] "
                    if not rationale_note.startswith(prefix):
                        rationale_note = prefix + (orig_note if orig_note else "Hoàn thiện Merchant Portal dựa trên API thực tế từ Backend.")

                row_data = [
                    global_seq,
                    t['task_id'],
                    block['block_id'],
                    subblock['sub_name'].split('. ')[1] if '. ' in subblock['sub_name'] else subblock['sub_name'],
                    f"W{t['orig_week']:02d}",
                    t['track'],
                    t['name'],
                    t['goal'],
                    t['how'],
                    t['output'],
                    t['criteria'],
                    t['test'],
                    t['est'],
                    t['priority'],
                    t['dependencies'],
                    t['status'],
                    t['evidence'],
                    rationale_note,
                    t['db_task_id'],
                    t['db_focus'],
                    t['db_output']
                ]

                for col_idx, val in enumerate(row_data, 1):
                    cell = ws.cell(row=current_row, column=col_idx, value=val)
                    cell.font = regular_font
                    cell.border = table_border
                    
                    # Alignments
                    if col_idx in [1, 2, 3, 5, 6, 13, 14, 16, 19]:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

                    # Bold font for priority and status
                    if col_idx in [14, 16]:
                        cell.font = bold_font

                ws.row_dimensions[current_row].height = 24
                current_row += 1
                global_seq += 1

    max_row = current_row - 1

    # 4. Data Validation (Dropdowns giống trang cũ)
    # Dropdown for Trạng thái (Column P)
    dv_status = DataValidation(type="list", formula1='"Todo,In Progress,Blocked,Review,Done,Deferred"', allow_blank=True)
    ws.add_data_validation(dv_status)
    dv_status.add(f"P3:P{max_row}")

    # Dropdown for Ưu tiên (Column N)
    dv_prio = DataValidation(type="list", formula1='"Must,Should,Stretch"', allow_blank=True)
    ws.add_data_validation(dv_prio)
    dv_prio.add(f"N3:N{max_row}")

    # 5. Conditional Formatting (Định dạng có điều kiện giống trang cũ)
    cf_range_p = f"P3:P{max_row}"
    fill_done = PatternFill(patternType='solid', fgColor='FFC6EFCE', bgColor='FFC6EFCE')
    fill_blocked = PatternFill(patternType='solid', fgColor='FFFFC7CE', bgColor='FFFFC7CE')
    fill_in_prog = PatternFill(patternType='solid', fgColor='FFFFEB9C', bgColor='FFFFEB9C')
    fill_review = PatternFill(patternType='solid', fgColor='FFD9EAF7', bgColor='FFD9EAF7')
    fill_deferred = PatternFill(patternType='solid', fgColor='FFE7E6E6', bgColor='FFE7E6E6')

    ws.conditional_formatting.add(cf_range_p, FormulaRule(formula=['$P3="Done"'], fill=fill_done))
    ws.conditional_formatting.add(cf_range_p, FormulaRule(formula=['$P3="Blocked"'], fill=fill_blocked))
    ws.conditional_formatting.add(cf_range_p, FormulaRule(formula=['$P3="In Progress"'], fill=fill_in_prog))
    ws.conditional_formatting.add(cf_range_p, FormulaRule(formula=['$P3="Review"'], fill=fill_review))
    ws.conditional_formatting.add(cf_range_p, FormulaRule(formula=['$P3="Deferred"'], fill=fill_deferred))

    cf_range_n = f"N3:N{max_row}"
    fill_must = PatternFill(patternType='solid', fgColor='FFF4CCCC', bgColor='FFF4CCCC')
    fill_should = PatternFill(patternType='solid', fgColor='FFFFF2CC', bgColor='FFFFF2CC')
    fill_stretch = PatternFill(patternType='solid', fgColor='FFD9EAD3', bgColor='FFD9EAD3')

    ws.conditional_formatting.add(cf_range_n, FormulaRule(formula=['$N3="Must"'], fill=fill_must))
    ws.conditional_formatting.add(cf_range_n, FormulaRule(formula=['$N3="Should"'], fill=fill_should))
    ws.conditional_formatting.add(cf_range_n, FormulaRule(formula=['$N3="Stretch"'], fill=fill_stretch))

    # Column dimensions
    col_widths = {
        'A': 9,   # STT
        'B': 13,  # Task ID
        'C': 10,  # Cục
        'D': 32,  # Phân nhóm
        'E': 10,  # Tuần gốc
        'F': 14,  # Track
        'G': 45,  # Task cần làm
        'H': 35,  # Mục tiêu học
        'I': 45,  # Cách thực hiện
        'J': 35,  # Đầu ra bắt buộc
        'K': 45,  # Acceptance criteria
        'L': 35,  # Test/kiểm chứng
        'M': 12,  # Giờ
        'N': 10,  # Ưu tiên
        'O': 14,  # Phụ thuộc
        'P': 16,  # Trạng thái
        'Q': 20,  # Link bằng chứng
        'R': 35,  # Ghi chú
        'S': 16,  # DB Task ID
        'T': 35,  # DB focus
        'U': 35   # DB output
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Freeze panes at C3
    ws.freeze_panes = 'C3'

    wb.save(wb_path)
    print(f"SUCCESS: Created sheet '{new_sheet_name}' with {global_seq - 1} tasks in {wb_path}!")

if __name__ == '__main__':
    build_block_plan()
