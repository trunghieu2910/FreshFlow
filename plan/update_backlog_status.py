import openpyxl

def update_status():
    file_path = 'd:/FreshFlow/plan/backlog-freshflow-mvp-12-tuan-updated.xlsx'
    completed_tasks = [
        'FF-03-05-1',
        'FF-03-05-2',
        'FF-03-06-1',
        'FF-03-06-2',
        'FF-03-07-1',
        'FF-03-07-2',
    ]
    try:
        wb = openpyxl.load_workbook(file_path)
        sheet = wb['Backlog 12 Tuan']
        updated_count = 0
        for r in range(1, 100):
            val = sheet.cell(row=r, column=1).value
            if val in completed_tasks:
                sheet.cell(row=r, column=16, value='Done')
                print(f'Row {r} ({val}) Col 16 updated to Done')
                updated_count += 1
        wb.save(file_path)
        print(f'SUCCESS: Updated {updated_count} tasks in backlog-freshflow-mvp-12-tuan-updated.xlsx!')
    except PermissionError:
        print('File is currently opened in Excel. Please close it and rerun.')

if __name__ == '__main__':
    update_status()
